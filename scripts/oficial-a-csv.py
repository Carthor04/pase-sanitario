"""Convierte los Excel oficiales del Instituto Oviedo (data/oficial/) a las plantillas CSV de
data/templates/ y los valida antes de importar. Escribe en data/oficial/csv/ (ignorada en git:
datos de menores). Luego se importan con scripts/import-csv.mjs.

    python scripts/oficial-a-csv.py

Entradas en data/oficial/:
  - "Listas 2026-27.xlsx": una pestaña por grupo (#, matrícula, nombre con apellidos primero).
  - "alumnos_sexo_para_verificar.xlsx": sexo H/M por alumno, verificado por la escuela.
  - "personal_para_llenar.xlsx": nombre, correo institucional y roles.
  - "Jornada maestr@s  26 -27.xlsx": horario POR DOCENTE (cada celda es el grupo).
  - "LS Tracker_Term 1.xlsx": listas de Inglés por grado y nivel.

Reglas del horario:
  - En Tecnología y robótica, "2ºAT" es Tecnología y "2ºA" es Robótica.
  - Inglés se da por grado y nivel, cada nivel en el salón de un grupo (ENGLISH_ROOMS). En esas
    horas el salón tiene la clase "Inglés básico/intermedio/avanzado" con su docente, y la lista
    del salón es la del nivel (secciones.csv).
  - Artes viene por grado ("1eros.") y aún sin docente: se carga en cada grupo del grado.
"""
import collections
import csv
import difflib
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'data' / 'oficial'
OUT = SRC / 'csv'
OUT.mkdir(exist_ok=True)

GROUPS = [f'{g}°{s}' for g in '123' for s in 'ABC']
ROOMS = {'2°A': 'Salón 1', '1°A': 'Salón 2', '2°B': 'Salón 3', '1°B': 'Salón 4', '2°C': 'Salón 5',
         '1°C': 'Salón 6', '3°C': 'Salón 7', '3°B': 'Salón 8', '3°A': 'Salón 9'}
# Nivel de Inglés → (salón donde se da, correo del docente). El salón es el de un grupo del grado.
ENGLISH_TEACHER = {
    'básico': 'carlosfrancom@institutooviedo.edu.mx',
    'intermedio': 'victoriaenriquezh@institutooviedo.edu.mx',
    'avanzado': 'luisramirezg@institutooviedo.edu.mx',
}
ENGLISH_ROOMS = {
    ('1', 'básico'): '1°C', ('1', 'intermedio'): '1°B', ('1', 'avanzado'): '1°A',
    ('2', 'básico'): '2°C', ('2', 'intermedio'): '2°B', ('2', 'avanzado'): '2°A',
    ('3', 'básico'): '3°C', ('3', 'intermedio'): '3°B', ('3', 'avanzado'): '3°A',
}
# Correos extra de una misma persona (la directora usa dos cuentas).
EXTRA_ACCOUNTS = {'gaby.lara@institutooviedo.edu.mx': ('Gaby Lara', 'director')}
PART = {'de', 'del', 'la', 'los', 'las', 'y'}
DAYS = {'lunes': 'lunes', 'martes': 'martes', 'miércoles': 'miércoles', 'miercoles': 'miércoles',
        'jueves': 'jueves', 'viernes': 'viernes'}
errors = []


def title(name):
    words = ' '.join(str(name).split()).rstrip('.').split(' ')
    return ' '.join(w.lower() if k and w.lower() in PART else w[:1].upper() + w[1:].lower()
                    for k, w in enumerate(words))


def norm(s):
    s = unicodedata.normalize('NFD', str(s)).encode('ascii', 'ignore').decode().upper()
    return ' '.join(re.sub(r'[^A-Z ]', ' ', s).split())


def write(name, header, rows):
    with open(OUT / name, 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


# ---------------------------------------------------------------- grupos
write('grupos.csv', ['grupo', 'salon', 'orden'], [[g, ROOMS[g], i + 1] for i, g in enumerate(GROUPS)])

# ---------------------------------------------------------------- alumnos
sexo = {}
for ws in openpyxl.load_workbook(SRC / 'alumnos_sexo_para_verificar.xlsx', data_only=True).worksheets:
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            sexo[(ws.title, int(r[0]))] = str(r[2] or '').strip().upper()
students = []
for ws in openpyxl.load_workbook(SRC / 'Listas 2026-27.xlsx', data_only=True).worksheets:
    g = ws.title.replace('º', '°').strip()
    seen = set()
    for r in ws.iter_rows(values_only=True):
        if isinstance(r[0], (int, float)) and r[2]:
            n = int(r[0])
            if n in seen:
                errors.append(('alumnos', g, f'número {n} repetido'))
            seen.add(n)
            s = sexo.get((g, n), '')
            if s not in ('H', 'M'):
                errors.append(('alumnos', g, f'#{n} sin sexo H/M'))
            students.append([g, n, title(r[2]), s])
write('alumnos.csv', ['grupo', 'numero_lista', 'nombre_completo', 'sexo'], students)

# ---------------------------------------------------------------- personal
staff = []
emails = set()
for r in openpyxl.load_workbook(SRC / 'personal_para_llenar.xlsx', data_only=True)['personal'].iter_rows(
        min_row=2, values_only=True):
    name, email, roles = (r + (None,) * 3)[:3]
    email = (email or '').strip().lower()
    if not email:
        if name:
            errors.append(('personal', name, 'sin correo: no podrá entrar ni se le asigna horario'))
        continue
    if not name and email in EXTRA_ACCOUNTS:
        name, roles = EXTRA_ACCOUNTS[email]
    if not name:
        errors.append(('personal', email, 'sin nombre ni roles: no se importa'))
        continue
    if not re.fullmatch(r'[^@\s]+@institutooviedo\.edu\.mx', email):
        errors.append(('personal', email, 'el correo no es @institutooviedo.edu.mx'))
    if email in emails:
        errors.append(('personal', email, 'correo repetido'))
    emails.add(email)
    staff.append([title(name), email, (roles or 'teacher').strip(), ''])
write('personal.csv', ['nombre_completo', 'correo', 'roles', 'whatsapp'], staff)

# ---------------------------------------------------------------- horario
# Pestaña del horario → correo del docente ('' = sin docente por ahora).
TEACHER = {
    'Matemáticas I y II': 'hectorfernandezh@institutooviedo.edu.mx',
    'Ignacio Franco': 'ignacio.franco@institutooviedo.edu.mx',
    'Cívica y ética': 'mariselasancheze@institutooviedo.edu.mx',
    'Español II,III': 'amalinallitecuapetlae@institutooviedo.edu.mx',
    'Español I, II': 'eduardorizos@institutooviedo.edu.mx',
    'Biología': 'edgarchavezb@institutooviedo.edu.mx',
    'Geografía y Quimica': 'andreacordobac@institutooviedo.edu.mx',
    'Ed. física': 'carlos.diaz@institutooviedo.edu.mx',
    'historia': 'alanmaganao@institutooviedo.edu.mx',
    'OYT': 'andreanoez@institutooviedo.edu.mx',
    'Tecnología y robótica': 'saelcarrillot@institutooviedo.edu.mx',
    'Emprendimiento 2º': 'mariagutierrezr@institutooviedo.edu.mx',
    'Emprendimiento 1º y 3º': 'jorge.esqueda@institutooviedo.edu.mx',
    'Artes': '',
    'Inglés': '',  # las tres pestañas de Inglés: el docente sale del nivel (ENGLISH_*)
}


def sheet_key(title_):
    t = title_.strip().rstrip('.')
    return 'Inglés' if t.startswith('Inglés') else t


def subject(sheet, group, tech):
    if sheet.startswith('Matemáticas') or sheet == 'Ignacio Franco':
        return 'Matemáticas'
    if sheet.startswith('Español'):
        return 'Español'
    if sheet.startswith('Geografía'):
        return 'Geografía' if group.startswith('1') else 'Química'
    if sheet.startswith('Tecnología'):
        return 'Tecnología' if tech else 'Robótica'
    if sheet.startswith('Emprendimiento'):
        return 'Emprendimiento'
    return {'Cívica y ética': 'Cívica y Ética', 'Biología': 'Biología', 'Ed. física': 'Educación Física',
            'historia': 'Historia', 'OYT': 'Orientación y Tutoría', 'Artes': 'Artes'}[sheet]


def hhmm(s):
    return [f'{int(h):02d}:{m}' for h, m in re.findall(r'(\d{1,2})\s*:\s*(\d{2})', s or '')]


schedule = set()
for ws in openpyxl.load_workbook(SRC / 'Jornada maestr@s  26 -27.xlsx').worksheets:
    key = sheet_key(ws.title)
    if key not in TEACHER:
        continue  # Baños (rol de guardias) y pestañas vacías
    email = TEACHER[key]
    if email and email not in emails:
        errors.append(('horario', ws.title, f'{email} no está en personal'))
    # Celdas combinadas (Inglés va de lunes a viernes en una sola celda).
    value = {}
    for m in ws.merged_cells.ranges:
        v = ws.cell(row=m.min_row, column=m.min_col).value
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                value[(r, c)] = v
    cell = lambda r, c: value.get((r, c), ws.cell(row=r, column=c).value)
    hdr = next(c for row in ws.iter_rows() for c in row if str(c.value or '').strip().lower() == 'hora')
    daycol = {c.column: DAYS[str(c.value).strip().lower()] for c in ws[hdr.row]
              if str(c.value or '').strip().lower() in DAYS}
    for r in range(hdr.row + 1, ws.max_row + 1):
        t = hhmm(str(ws.cell(row=r, column=1).value or ''))
        if len(t) != 2:
            continue
        for col, day in daycol.items():
            raw = re.sub(r'\s+', '', str(cell(r, col) or ''))
            if not raw or raw.startswith('RE') or raw == 'R':
                continue
            m = re.fullmatch(r'([123])[º°o]([ABC])(T?)', raw)
            grade = re.search(r'([123])', raw)
            if m:
                g = f'{m.group(1)}°{m.group(2)}'
                schedule.add((g, day, t[0], t[1], subject(key, g, bool(m.group(3))), email))
            elif grade and re.search(r'\d(EROS|ROS|DOS)\.?$', raw.upper()):
                gr = grade.group(1)
                if key == 'Inglés':
                    for level, room in ((lv, ENGLISH_ROOMS[(gr, lv)]) for lv in ENGLISH_TEACHER):
                        schedule.add((room, day, t[0], t[1], f'Inglés {level}', ENGLISH_TEACHER[level]))
                else:
                    for s in 'ABC':
                        schedule.add((f'{gr}°{s}', day, t[0], t[1], subject(key, f'{gr}°{s}', False), email))
            else:
                errors.append(('horario', ws.title, f'celda no reconocida: {raw!r}'))

schedule = sorted(schedule)
slot = collections.defaultdict(list)
for e in schedule:
    slot[(e[0], e[1], e[2])].append(e)
for k, v in slot.items():
    if len(v) > 1:
        errors.append(('horario', f'{k[0]} {k[1]} {k[2]}', 'empalme: ' + ' / '.join(f'{x[4]} {x[5]}' for x in v)))
order = {d: i for i, d in enumerate(['lunes', 'martes', 'miércoles', 'jueves', 'viernes'])}
schedule.sort(key=lambda e: (GROUPS.index(e[0]), order[e[1]], e[2]))
write('horario.csv', ['grupo', 'dia', 'inicio', 'fin', 'materia', 'correo_maestro'], schedule)

# ---------------------------------------------------------------- Inglés por nivel
LEVEL = {'BASIC': 'básico', 'INTERMEDIATE': 'intermedio', 'ADVANCED': 'avanzado'}
by_grade = collections.defaultdict(list)
for s in students:
    by_grade[s[0][0]].append(s)


def score(raw, s):
    a, b = norm(raw), norm(s[2])
    if set(a.split()) <= set(b.split()):  # mismo alumno escrito con nombre primero o incompleto
        return 0.99
    return difflib.SequenceMatcher(None, a, b).ratio()


members = []
placed = collections.defaultdict(list)
for ws in openpyxl.load_workbook(SRC / 'LS Tracker_Term 1.xlsx', data_only=True).worksheets:
    meta, names = {}, []
    for r in ws.iter_rows(min_row=1, max_row=200, max_col=3, values_only=True):
        if isinstance(r[0], str) and r[0].strip().rstrip(':') in ('Grade', 'Level'):
            meta[r[0].strip().rstrip(':')] = str(r[1]).strip()
        elif isinstance(r[0], (int, float)) and isinstance(r[1], str) and r[1].strip():
            names.append(r[1].strip())
    gr, level = meta['Grade'][0], LEVEL[meta['Level'].upper()]
    room = ENGLISH_ROOMS[(gr, level)]
    for raw in names:
        ranked = sorted(by_grade[gr], key=lambda s: -score(raw, s))
        best, second = score(raw, ranked[0]), score(raw, ranked[1])
        if best < 0.86 or best - second < 0.05:
            errors.append(('inglés', ws.title, f'{raw!r} ≈ {ranked[0][2]} ({ranked[0][0]}) — revisar'))
        s = ranked[0]
        placed[(s[0], s[1])].append(ws.title)
        members.append([room, f'Inglés {level}', s[0], s[1]])
for k, v in placed.items():
    if len(v) > 1:
        errors.append(('inglés', f'{k[0]} #{k[1]}', 'aparece en más de una lista: ' + ', '.join(v)))
for s in students:
    if (s[0], s[1]) not in placed:
        errors.append(('inglés', f'{s[0]} #{s[1]}', f'{s[2]} no aparece en ninguna lista de Inglés'))
write('secciones.csv', ['salon_grupo', 'materia', 'grupo', 'numero_lista'], members)

# ---------------------------------------------------------------- guardias del baño
# Hoja "Baños" del horario: solo trae el nombre de pila. FRANCO es Ignacio Franco (Carlos
# Franco tiene clase en esas horas), ANDREA es Andrea Córdoba (Andrea Noe tiene clase el
# miércoles 12:15), CARLOS es Carlos Franco y LUCY es Luz del Carmen Tejeda (2026-10-23).
# PATRICK aún no está en el personal: su guardia queda sin asignar.
GUARD = {
    'EDUARDO': 'eduardorizos@institutooviedo.edu.mx',
    'ANDRES': 'saelcarrillot@institutooviedo.edu.mx',
    'FRANCO': 'ignacio.franco@institutooviedo.edu.mx',
    'ALAN': 'alanmaganao@institutooviedo.edu.mx',
    'LUCY': 'coordinacion.secundaria@institutooviedo.edu.mx',
    'HECTOR': 'hectorfernandezh@institutooviedo.edu.mx',
    'MARISELA': 'mariselasancheze@institutooviedo.edu.mx',
    'ANDREA': 'andreacordobac@institutooviedo.edu.mx',
    'EDGAR': 'edgarchavezb@institutooviedo.edu.mx',
    'AMALINALLI': 'amalinallitecuapetlae@institutooviedo.edu.mx',
    'RICARDO': 'luisramirezg@institutooviedo.edu.mx',
    'VICTORIA': 'victoriaenriquezh@institutooviedo.edu.mx',
    'CARLOS': 'carlosfrancom@institutooviedo.edu.mx',
    'PATRICK': '',
}
guards = []
ws = next(w for w in openpyxl.load_workbook(SRC / 'Jornada maestr@s  26 -27.xlsx').worksheets
          if norm(w.title) == 'BANOS')
hdr = next(c for row in ws.iter_rows() for c in row if str(c.value or '').strip().lower() == 'hora')
daycol = {c.column: DAYS[str(c.value).strip().lower()] for c in ws[hdr.row]
          if str(c.value or '').strip().lower() in DAYS}
for r in range(hdr.row + 1, ws.max_row + 1):
    t = hhmm(str(ws.cell(row=r, column=1).value or ''))
    if len(t) != 2:
        continue
    for col, day in daycol.items():
        name = norm(ws.cell(row=r, column=col).value or '')
        if not name or name.startswith('R E'):
            continue  # receso
        if name not in GUARD:
            errors.append(('guardias', f'{day} {t[0]}', f'nombre no reconocido: {name}'))
            continue
        if not GUARD[name]:
            errors.append(('guardias', f'{day} {t[0]}', f'{name} no está en el personal: queda sin asignar'))
        elif GUARD[name] not in emails:
            errors.append(('guardias', f'{day} {t[0]}', f'{GUARD[name]} no está en personal'))
        guards.append([day, t[0], t[1], GUARD[name]])
write('guardias.csv', ['dia', 'inicio', 'fin', 'correo_maestro'], guards)

# ---------------------------------------------------------------- resumen
per = collections.Counter(s[0] for s in students)
hrs = collections.Counter(e[0] for e in schedule)
sin = collections.Counter(e[0] for e in schedule if not e[5])
print('Grupo  salón     alumnos  H   M   clases  sin docente')
for g in GROUPS:
    h = sum(1 for s in students if s[0] == g and s[3] == 'H')
    print(f'{g:5}  {ROOMS[g]:8}  {per[g]:7}  {h:2}  {per[g] - h:2}  {hrs[g]:6}  {sin[g]:4}')
lv = collections.Counter((m[0], m[1]) for m in members)
print('Inglés:', ', '.join(f'{k[1]} en {k[0]} ({ROOMS[k[0]]}): {v}' for k, v in sorted(lv.items())))
print(f'Total: {len(students)} alumnos, {len(staff)} personal, {len(schedule)} clases, {len(members)} en Inglés, '
      f'{len(guards)} guardias')
print('\nAvisos:' if errors else '\nSin avisos.')
for e in errors:
    print('  ', ' · '.join(map(str, e)))
